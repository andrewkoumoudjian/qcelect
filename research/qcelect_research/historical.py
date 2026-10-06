"""Province-wide Élections Québec historical result normalization.

This module is offline-only. It downloads the official polling-station archives,
parses each 2014+ riding file, validates vote reconciliation, and emits one
canonical long table suitable for boundary transposition and historical replay.
"""

from __future__ import annotations

import hashlib
import csv
import io
import json
import math
import re
import urllib.request
import zipfile
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from typing import Iterable

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from .elections_quebec import (
    GENERAL_ELECTION_ARCHIVES,
    _decode_bytes,
    _rows_from_text,
    _strip_trailing_garbage,
    archive_url,
)

MODERN_ELECTIONS = ("2014-04-07", "2018-10-01", "2022-10-03")

AGGREGATE_MARKERS = (
    "TOTAL DU SECTEUR",
    "GRAND TOTAL",
    "MAJORITÉ EN FAVEUR",
    "MAJORITE EN FAVEUR",
)

SOURCE_NOTE_MARKERS = (
    "LES RÉSULTATS",
    "LES RESULTATS",
)

SPECIAL_VOTE_MARKERS = (
    "VOTE DES DÉTENUS",
    "VOTE DES DETENUS",
    "VOTE HORS QUÉBEC",
    "VOTE HORS QUEBEC",
    "ÉLECTEURS HORS",
    "ELECTEURS HORS",
)

PARTY_ALIASES = {
    "P.Q.": "PQ",
    "Q.S.": "QS",
    "P.L.Q./Q.L.P.": "PLQ",
    "C.A.Q.-É.F.L.": "CAQ",
    "C.A.Q.-E.F.L.": "CAQ",
    "P.C.Q./C.P.Q.": "PCOQ",
    "É.A.P. - P.C.Q.": "PCOQ",
    "E.A.P. - P.C.Q.": "PCOQ",
    "O.N. - P.I.Q.": "ON",
    "P.V.Q./G.P.Q.": "PVQ",
    "P.M.L.Q.": "PMLQ",
    "P.C.Q-E.E.D.": "PCOQ",
}

# The Canadian Party reused a conservative abbreviation in 2022. Party
# identity must follow the election-specific source, not punctuation alone.
ELECTION_PARTY_ALIASES = {"2022-10-03": {"P.C.Q./C.P.Q": "PCANQ"}}

CANONICAL_COLUMNS = [
    "election",
    "riding_code",
    "riding",
    "polling_section",
    "polling_section_group",
    "municipality",
    "candidate",
    "party",
    "party_raw",
    "votes",
    "valid_votes",
    "rejected_votes",
    "registered_electors",
    "turnout",
    "is_merged_poll",
    "is_special_vote",
    "source_file",
]


@dataclass(frozen=True)
class HistoricalDataset:
    rows: pd.DataFrame
    diagnostics: dict[str, object]


def _contains_marker(row: Iterable[object], markers: Iterable[str]) -> bool:
    haystack = " | ".join(str(value) for value in row).upper()
    return any(marker in haystack for marker in markers)


def _clean_identifier(value: object) -> str:
    text = str(value).strip()
    if not text or text.lower() == "nan":
        return ""
    if re.fullmatch(r"-?\d+\.0+", text):
        return text.split(".", 1)[0]
    return text


def _number(value: object) -> float:
    if str(value).strip().lower() in {"", "nan"}:
        return float("nan")
    try:
        parsed = float(str(value).strip())
    except ValueError as exc:
        raise ValueError(f"invalid official vote/elector count: {value!r}") from exc
    if not math.isfinite(parsed) or parsed < 0 or not parsed.is_integer():
        raise ValueError(f"invalid official vote/elector count: {value!r}")
    return parsed


def _split_candidate_header(label: str, election_date: str) -> tuple[str, str, str]:
    clean = label.strip()
    tokens = clean.split()
    if len(tokens) < 2:
        return clean, "", ""

    # DGEQ party abbreviations are embedded at the end of the candidate label.
    # Some 2014 labels are multi-token coalitions such as
    # "É.A.P. - P.C.Q.", so rsplit(" ", 1) is not sufficient.
    aliases = {**PARTY_ALIASES, **ELECTION_PARTY_ALIASES.get(election_date, {})}
    party_raw = next(
        (
            alias
            for alias in sorted(aliases, key=len, reverse=True)
            if clean.endswith(f" {alias}")
        ),
        None,
    )
    if party_raw is None:
        candidate, party_raw = clean.rsplit(" ", 1)
    else:
        candidate = clean[: -(len(party_raw) + 1)]

    party_raw = party_raw.strip()
    return candidate.strip(), aliases.get(party_raw, party_raw), party_raw


def _header_and_records(
    raw: bytes, *, election_date: str = "", source_file: str = ""
) -> tuple[list[str], list[list[str]]]:
    rows = _rows_from_text(_decode_bytes(raw))
    if len(rows) < 2:
        raise ValueError("DGEQ file has fewer than two rows")

    header_index = next(
        (
            index
            for index, row in enumerate(rows[:25])
            if row and row[0].strip().rstrip(":").lower() == "header"
        ),
        None,
    )

    if header_index is not None:
        headers = [
            value.strip() for value in _strip_trailing_garbage(rows[header_index][1:])
        ]
        record_start = header_index + 1
    else:
        # The current official 2014 archive differs from the later 2018/2022
        # exports and from older copied research files: each riding CSV starts
        # directly with the named header row, with no Election:/Header: prefix.
        direct_headers = [value.strip() for value in _strip_trailing_garbage(rows[0])]
        if (
            election_date == "2022-10-03"
            and PurePosixPath(source_file).name.lower()
            == "dge-80.10_acadie_sans_se.csv"
            and direct_headers[0] == "S"
        ):
            # The official Acadie 2022 CSV labels its riding-code column S.
            direct_headers[0] = "Code"
        structural = {"Code", "Circonscription", "S.V.", "É.I.", "B.V.", "B.R."}
        if not structural.issubset(direct_headers):
            sample = [row[:5] for row in rows[:8]]
            raise ValueError(
                "DGEQ file has neither a Header: prefix nor a recognized "
                f"direct header; first rows={sample!r}"
            )
        headers = direct_headers
        record_start = 1
    if not headers:
        raise ValueError("modern DGEQ file has an empty header")

    records: list[list[str]] = []
    for raw_row in rows[record_start:]:
        row = _strip_trailing_garbage(raw_row)
        if not row:
            continue
        if row[0].strip().rstrip(":").lower() in {"parties", "party"}:
            continue
        if _contains_marker(row, (*AGGREGATE_MARKERS, *SOURCE_NOTE_MARKERS)):
            continue

        is_special = _contains_marker(row, SPECIAL_VOTE_MARKERS)
        if len(row) < len(headers):
            if not is_special:
                raise ValueError(
                    f"record has {len(row)} fields but header has {len(headers)}"
                )
            row = [*row, *([""] * (len(headers) - len(row)))]
        if len(row) > len(headers):
            raise ValueError(
                f"record has {len(row)} fields but header has {len(headers)}"
            )
        records.append(row)

    return headers, records


def parse_modern_riding_csv(
    raw: bytes,
    *,
    election_date: str,
    source_file: str,
) -> pd.DataFrame:
    """Parse one 2014/2018/2022 per-riding file into canonical long rows."""

    if election_date not in MODERN_ELECTIONS:
        raise ValueError(f"{election_date} is not a supported modern election")

    headers, records = _header_and_records(
        raw, election_date=election_date, source_file=source_file
    )
    required = {
        "Code",
        "Circonscription",
        "Nom des Municipalités",
        "S.V.",
        "É.I.",
        "B.V.",
        "B.R.",
    }
    missing = required.difference(headers)
    if missing:
        raise ValueError(f"modern DGEQ header missing columns: {sorted(missing)}")

    electors_index = headers.index("É.I.")
    valid_index = headers.index("B.V.")
    if valid_index <= electors_index + 1:
        raise ValueError("no candidate columns between É.I. and B.V.")

    candidate_columns = headers[electors_index + 1 : valid_index]
    parsed_candidates = [
        _split_candidate_header(label, election_date) for label in candidate_columns
    ]
    if any(not party_raw for _, _, party_raw in parsed_candidates):
        raise ValueError("candidate header missing party suffix")

    output: list[dict[str, object]] = []
    declared_totals: list[list[int]] = []

    for record in records:
        row = dict(zip(headers, record, strict=True))
        is_special = _contains_marker(record, SPECIAL_VOTE_MARKERS)
        polling_section = _clean_identifier(row["S.V."])
        municipality = str(row["Nom des Municipalités"]).strip()
        if not polling_section and (
            re.match(r"^BVA\s*\d+(?:\s*:|$)", municipality, re.IGNORECASE)
            or municipality.lower().startswith("vote ")
        ):
            is_special = True
            polling_section = f"special:{municipality}"
        polling_section_group = str(row.get("Regroupement", "")).strip()
        registered = _number(row["É.I."])
        valid = _number(row["B.V."])
        rejected = _number(row["B.R."])

        candidate_votes = [_number(row[column]) for column in candidate_columns]
        candidate_votes = [
            0.0 if np.isnan(value) else value for value in candidate_votes
        ]

        if municipality.upper() == "TOTAL DE LA CIRCONSCRIPTION":
            declared_totals.append([int(value) for value in candidate_votes])
            continue
        if not polling_section:
            raise ValueError(
                f"{source_file}: unrecognized row without section: {municipality!r}"
            )

        if not is_special or not np.isnan(valid):
            if np.isnan(valid):
                raise ValueError(f"{source_file}: regular polling row has no B.V.")
            difference = abs(sum(candidate_votes) - valid)
            if difference > 1e-9:
                raise ValueError(
                    f"{source_file}: candidate votes do not reconcile to B.V.; "
                    f"poll={polling_section!r}, diff={difference}"
                )

        turnout = (
            valid / registered
            if not np.isnan(valid) and not np.isnan(registered) and registered > 0
            else float("nan")
        )
        try:
            poll_number = int(polling_section)
        except ValueError:
            poll_number = -1

        for column, (candidate, party, party_raw), votes in zip(
            candidate_columns,
            parsed_candidates,
            candidate_votes,
            strict=True,
        ):
            output.append(
                {
                    "election": election_date,
                    "riding_code": _clean_identifier(row["Code"]),
                    "riding": str(row["Circonscription"]).strip(),
                    "polling_section": polling_section,
                    "polling_section_group": polling_section_group,
                    "municipality": str(row["Nom des Municipalités"]).strip(),
                    "candidate": candidate,
                    "party": party,
                    "party_raw": party_raw,
                    "votes": int(votes),
                    "valid_votes": int(valid) if not np.isnan(valid) else pd.NA,
                    "rejected_votes": int(rejected)
                    if not np.isnan(rejected)
                    else pd.NA,
                    "registered_electors": int(registered)
                    if not np.isnan(registered)
                    else pd.NA,
                    "turnout": turnout,
                    "is_merged_poll": poll_number >= 500 or bool(polling_section_group),
                    "is_special_vote": is_special,
                    "source_file": source_file,
                }
            )

    frame = pd.DataFrame(output, columns=CANONICAL_COLUMNS)
    if frame.empty:
        raise ValueError(f"{source_file}: no polling-station records parsed")
    if declared_totals:
        if len(declared_totals) != 1:
            raise ValueError(f"{source_file}: multiple riding totals")
        actual = [
            int(
                frame.loc[
                    (frame["candidate"] == candidate) & (frame["party"] == party),
                    "votes",
                ].sum()
            )
            for candidate, party, _ in parsed_candidates
        ]
        if actual != declared_totals[0]:
            raise ValueError(
                f"{source_file}: candidate sums differ from declared riding totals"
            )
    frame.attrs["declared_totals_checked"] = bool(declared_totals)
    return frame


def _validate_poll_uniqueness(frame: pd.DataFrame) -> None:
    regular = frame.copy()

    # In long format a polling-section ID legitimately repeats once per
    # candidate. A repeated candidate within the same poll means the poll row
    # itself was ingested more than once.
    candidate_keys = [
        "election",
        "riding_code",
        "polling_section",
        "candidate",
        "party",
    ]
    duplicate_candidates = regular.duplicated(candidate_keys, keep=False)
    if duplicate_candidates.any():
        sample = regular.loc[duplicate_candidates, candidate_keys].head()
        raise ValueError(
            "duplicated polling-section candidate rows within election/riding: "
            f"{sample.to_dict(orient='records')}"
        )

    # The same polling-section ID must not be sourced from two different
    # per-riding files either.
    poll_sources = regular[
        ["election", "riding_code", "polling_section", "source_file"]
    ].drop_duplicates()
    source_counts = poll_sources.groupby(
        ["election", "riding_code", "polling_section"], dropna=False
    )["source_file"].nunique()
    duplicates = source_counts[source_counts > 1]
    if not duplicates.empty:
        sample = duplicates.head().index.tolist()
        raise ValueError(
            f"duplicated polling-section IDs within election/riding: {sample}"
        )


def validate_historical_long(frame: pd.DataFrame) -> dict[str, object]:
    """Validate canonical historical rows and return reproducible diagnostics."""

    missing = set(CANONICAL_COLUMNS).difference(frame.columns)
    if missing:
        raise ValueError(f"historical frame missing columns: {sorted(missing)}")
    if frame.empty:
        raise ValueError("historical frame is empty")

    _validate_poll_uniqueness(frame)

    poll_keys = [
        "election",
        "riding_code",
        "polling_section",
        "source_file",
        "is_special_vote",
    ]
    totals = (
        frame.groupby(poll_keys, dropna=False, sort=False)
        .agg(candidate_votes=("votes", "sum"), valid_votes=("valid_votes", "first"))
        .reset_index()
    )
    regular = totals[totals["valid_votes"].notna()]
    mismatch = (regular["candidate_votes"] - regular["valid_votes"]).abs()
    if not mismatch.empty and float(mismatch.max()) > 1e-9:
        raise ValueError(
            "historical candidate votes do not reconcile to valid votes; "
            f"max diff={float(mismatch.max())}"
        )

    regular_rows = frame[~frame["is_special_vote"]]
    riding_count = int(
        regular_rows[["election", "riding_code"]].drop_duplicates().shape[0]
    )
    poll_count = int(
        regular_rows[["election", "riding_code", "polling_section"]]
        .drop_duplicates()
        .shape[0]
    )

    party_votes = (
        frame.groupby(["election", "party"], dropna=False)["votes"]
        .sum()
        .astype(int)
        .reset_index()
    )

    return {
        "elections": sorted(str(value) for value in frame["election"].unique()),
        "riding_election_count": riding_count,
        "regular_poll_count": poll_count,
        "special_poll_count": int(
            frame.loc[
                frame["is_special_vote"],
                ["election", "riding_code", "polling_section", "source_file"],
            ]
            .drop_duplicates()
            .shape[0]
        ),
        "merged_poll_count": int(
            regular_rows.loc[
                regular_rows["is_merged_poll"],
                ["election", "riding_code", "polling_section"],
            ]
            .drop_duplicates()
            .shape[0]
        ),
        "party_votes": party_votes.to_dict(orient="records"),
    }


def _is_riding_result_member(name: str, election_date: str) -> bool:
    """Select only official per-riding result CSVs from a DGEQ archive."""

    basename = PurePosixPath(name).name.lower()
    if not basename.endswith((".csv", ".xlsx")):
        return False
    if election_date == "2014-04-07":
        return basename.endswith("_officiels2014.csv")
    if election_date in {"2018-10-01", "2022-10-03"}:
        return bool(re.fullmatch(r"dge-80[.-]10_.+_sans[-_]se\.(?:csv|xlsx)", basename))
    return False


def parse_election_archive(
    raw_zip: bytes,
    *,
    election_date: str,
    expected_riding_count: int | None = None,
    require_declared_totals: bool = False,
) -> HistoricalDataset:
    """Parse all modern riding CSVs from one official DGEQ ZIP archive."""

    if election_date not in MODERN_ELECTIONS:
        raise ValueError(f"{election_date} is not a supported modern election")

    frames: list[pd.DataFrame] = []
    ignored_csv_members: list[str] = []
    with zipfile.ZipFile(io.BytesIO(raw_zip)) as archive:
        csv_members = sorted(
            name
            for name in archive.namelist()
            if name.lower().endswith((".csv", ".xlsx")) and not name.endswith("/")
        )
        members = [
            name
            for name in csv_members
            if _is_riding_result_member(name, election_date)
        ]
        ignored_csv_members = [name for name in csv_members if name not in set(members)]
        if not members:
            raise ValueError(
                f"{election_date}: DGEQ archive contains no recognized "
                "per-riding result CSVs"
            )

        for member in members:
            try:
                raw = archive.read(member)
                if member.lower().endswith(".xlsx"):
                    # Four 2018 riding files are official spreadsheets. Adapt
                    # their tabular transport into the same strict row parser.
                    workbook = load_workbook(
                        io.BytesIO(raw), read_only=True, data_only=True
                    )
                    try:
                        if len(workbook.worksheets) != 1:
                            raise ValueError("expected one official results worksheet")
                        text = io.StringIO()
                        writer = csv.writer(text)
                        writer.writerows(
                            workbook.worksheets[0].iter_rows(values_only=True)
                        )
                        raw = text.getvalue().encode("utf-8")
                    finally:
                        workbook.close()
                frame = parse_modern_riding_csv(
                    raw,
                    election_date=election_date,
                    source_file=member,
                )
            except Exception as exc:
                raise ValueError(
                    f"{election_date}: failed parsing archive member {member!r}"
                ) from exc
            frames.append(frame)

    rows = pd.concat(frames, ignore_index=True)
    if require_declared_totals and any(not frame.attrs["declared_totals_checked"] for frame in frames):
        raise ValueError(f"{election_date}: missing declared riding candidate totals")
    diagnostics = validate_historical_long(rows)
    if expected_riding_count is not None and (
        len(frames) != expected_riding_count
        or rows["riding_code"].nunique() != expected_riding_count
        or any(frame["riding_code"].nunique() != 1 for frame in frames)
    ):
        raise ValueError(
            f"{election_date}: expected exactly {expected_riding_count} riding files and IDs"
        )
    diagnostics.update(
        {
            "archive_url": archive_url(election_date),
            "archive_sha256": hashlib.sha256(raw_zip).hexdigest(),
            "source_file_count": len(frames),
            "source_files": members,
            "ignored_csv_members": ignored_csv_members,
        }
    )
    return HistoricalDataset(rows=rows, diagnostics=diagnostics)


def download_archive(
    election_date: str,
    cache_dir: str | Path,
    *,
    refresh: bool = False,
) -> Path:
    """Download one official archive once, retaining a local research cache."""

    if election_date not in GENERAL_ELECTION_ARCHIVES:
        raise KeyError(f"unknown election date {election_date}")

    cache = Path(cache_dir)
    cache.mkdir(parents=True, exist_ok=True)
    target = cache / f"{election_date}-resultats-bureau-vote.zip"
    if target.exists() and not refresh:
        return target

    request = urllib.request.Request(
        archive_url(election_date),
        headers={
            "Accept": "application/zip, application/octet-stream",
            "User-Agent": "qcelect-research/1.0 (+https://github.com/andrewkoumoudjian/qcelect)",
        },
    )
    with urllib.request.urlopen(request, timeout=60) as response:
        body = response.read()

    # Validate before replacing the cache entry.
    with zipfile.ZipFile(io.BytesIO(body)) as archive:
        if not archive.namelist():
            raise ValueError(f"{election_date}: downloaded archive is empty")

    temporary = target.with_suffix(".tmp")
    temporary.write_bytes(body)
    temporary.replace(target)
    return target


def build_historical_dataset(
    *,
    elections: Iterable[str] = MODERN_ELECTIONS,
    cache_dir: str | Path,
    refresh: bool = False,
) -> HistoricalDataset:
    """Download, normalize and validate all requested modern elections."""

    datasets: list[HistoricalDataset] = []
    for election_date in elections:
        path = download_archive(election_date, cache_dir, refresh=refresh)
        datasets.append(
            parse_election_archive(
                path.read_bytes(),
                election_date=election_date,
                expected_riding_count=125,
                require_declared_totals=True,
            )
        )

    rows = pd.concat([dataset.rows for dataset in datasets], ignore_index=True)
    diagnostics = validate_historical_long(rows)
    diagnostics["archives"] = [dataset.diagnostics for dataset in datasets]
    return HistoricalDataset(rows=rows, diagnostics=diagnostics)


def write_historical_dataset(
    dataset: HistoricalDataset,
    *,
    output_csv: str | Path,
    diagnostics_json: str | Path,
) -> None:
    output = Path(output_csv)
    diagnostics = Path(diagnostics_json)
    output.parent.mkdir(parents=True, exist_ok=True)
    diagnostics.parent.mkdir(parents=True, exist_ok=True)

    dataset.rows.to_csv(output, index=False, compression="gzip")
    diagnostics.write_text(
        json.dumps(dataset.diagnostics, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
